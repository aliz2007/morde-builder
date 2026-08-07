#!/usr/bin/env python3
"""Regenerate data/ and assets/img/ from the source Mordekaiser spreadsheet.

    pip install openpyxl
    python3 tools/build_data.py path/to/Mordekaiser_Matchup_Spreadsheet.xlsx

Reads the workbook, extracts all six sheets plus the embedded icon media and
every hyperlink, and writes data/matchups.json, data/guides.json and
assets/img/. The author's prose is copied verbatim — nothing here rewrites,
summarises, reorders or invents content.
"""
import openpyxl, re, json, os, shutil, sys, zipfile, tempfile

if len(sys.argv) < 2:
    sys.exit("usage: build_data.py <workbook.xlsx> [output_dir]")
SRC = sys.argv[1]
OUT = os.path.abspath(sys.argv[2] if len(sys.argv) > 2
                      else os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))
UNZ = tempfile.mkdtemp(prefix="morde-xlsx-")
with zipfile.ZipFile(SRC) as z:
    z.extractall(UNZ)

MATCHUPS_SHEET = "Reformatted Spreadsheet"


# ---------------------------------------------------------------- media
def drawing_map(idx):
    """(row, col) -> media filename, for the drawing attached to sheet `idx`."""
    dp = f"{UNZ}/xl/drawings/drawing{idx}.xml"
    rp = f"{UNZ}/xl/drawings/_rels/drawing{idx}.xml.rels"
    if not os.path.exists(dp):
        return {}
    x = open(dp, encoding="utf8").read()
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"',
                           open(rp, encoding="utf8").read()))
    out = {}
    for a in re.findall(r'<xdr:(?:two|one)CellAnchor.*?</xdr:(?:two|one)CellAnchor>', x, re.S):
        f = re.search(r'<xdr:from>\s*<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', a, re.S)
        e = re.search(r'r:embed="(rId\d+)"', a)
        if f and e:
            out[(int(f.group(2)) + 1, int(f.group(1)) + 1)] = rels.get(e.group(1))
    return out


# ---------------------------------------------------------------- workbook
# Not read_only: hyperlinks are only exposed on a fully loaded workbook.
wb = openpyxl.load_workbook(SRC, data_only=True)

LINKS = {}   # (sheet, row, col) -> url
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink is not None and c.hyperlink.target:
                LINKS[(ws.title, c.row, c.column)] = c.hyperlink.target


def grid(sheet, maxr, maxc):
    ws = wb[sheet]
    return [[c.value for c in row]
            for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc)]


def cell(rows, r, c):
    v = rows[r - 1][c - 1] if 0 < r <= len(rows) and c <= len(rows[r - 1]) else None
    return str(v).strip() if v not in (None, "") else ""


def link(sheet, r, c):
    return LINKS.get((sheet, r, c)) or ""


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')


# Official capitalisation, so search and display both behave. The author's own
# spelling is preserved on each record as `sourceName`.
CANON = {"Ganglank": "Gangplank", "Cho'gath": "Cho'Gath", "Bel'veth": "Bel'Veth",
         "K'sante": "K'Sante", "Kai'sa": "Kai'Sa", "Rek'sai": "Rek'Sai",
         "Vel'koz": "Vel'Koz", "Kog'maw": "Kog'Maw", "Leblanc": "LeBlanc"}
ALIAS = {"Gangplank": ["gp", "ganglank"], "Cho'Gath": ["chogath", "cho"],
         "Bel'Veth": ["belveth"], "K'Sante": ["ksante"], "Kai'Sa": ["kaisa"],
         "Rek'Sai": ["reksai"], "Vel'Koz": ["velkoz"], "Kog'Maw": ["kogmaw"],
         "LeBlanc": ["leblanc", "lb"], "Dr. Mundo": ["mundo", "drmundo"],
         "Jarvan IV": ["jarvan", "j4"], "Aurelion Sol": ["asol", "aurelionsol"],
         "Master Yi": ["yi"], "Twisted Fate": ["tf"], "Xin Zhao": ["xin"],
         "Tahm Kench": ["tahm", "kench"], "Lee Sin": ["lee"], "Mordekaiser": ["morde"],
         "Tryndamere": ["trynd", "tryn"], "Wukong": ["monkeyking"]}

# Icon identities are asserted ONLY where the workbook's own text confirms them:
# each Itemization Guide icon paired with the prose beside it (which names the
# item), plus Doran's Helm, which the build text names on the two matchups that
# use it. Everything else renders with a generic "Item n" alt.
#
# Slot 1 is deliberately absent. It holds a start item the author never names —
# across the 127 builds that share one icon, the author's own step 1 reads
# Rocketbelt, Dusk & Dawn, Plated Steelcaps, Riftmaker and so on, i.e. the text
# list begins at slot 2. An earlier version guessed "Bramble Vest" here and was
# wrong: Bramble Vest is image72, which the Itemization Guide names directly and
# which sits at slot 2 on Darius, whose step 1 is "Bramble Vest".
ITEM_NAMES = {
    "image116.jpg": "Doran's Helm",
    "image72.png": "Bramble Vest",
    "image108.png": "Hextech Rocketbelt", "image25.png": "Rylai's Crystal Scepter",
    "image149.png": "Cosmic Drive", "image144.png": "Rylai's Crystal Scepter",
    "image113.png": "Experimental Hexplate", "image134.png": "Riftmaker",
    "image115.png": "Dusk & Dawn", "image226.png": "Liandry's Torment",
    "image165.png": "Bloodletter's Curse", "image263.png": "Zhonya's Hourglass",
    "image69.png": "Kaenic Rookern", "image46.png": "Sterak's Gage",
    "image70.png": "Death's Dance", "image77.png": "Randuin's Omen",
    "image188.png": "Spirit Visage", "image73.png": "Force of Nature",
    "image64.png": "Abyssal Mask", "image109.png": "Frozen Heart",
    "image93.png": "Jak'Sho, The Protean", "image82.png": "Unending Despair",
    "image96.png": "Dead Man's Plate", "image127.png": "Dark Seal",
    "image88.png": "Mejai's Soulstealer", "image146.png": "Oblivion Orb",
    "image105.png": "Verdant Barrier", "image104.png": "Guardian Angel",
    "image100.png": "Serpent's Fang", "image63.jpg": "Plated Steelcaps",
    "image114.jpg": "Mercury's Treads", "image280.png": "Gluttonous Greaves",
    "image218.jpg": "Boots of Swiftness",
}
# Tiles that are literally text graphics in the workbook, not items.
PLACEHOLDER = {"image99.png": "Build Variety Items", "image143.png": "Flex Boots"}
SUMMONERS = {"image78.jpg": "Flash", "image94.jpg": "Ignite",
             "image118.jpg": "Ghost", "image204.jpg": "Teleport"}

used_media = set()

# ---------------------------------------------------------------- matchups
rows = grid(MATCHUPS_SHEET, 1124, 27)
IMG = drawing_map(2)


def g(i, c):
    return cell(rows, i + 1, c)


def rating(s):
    m = re.search(r'(\d+)\s*/\s*5', s)
    return int(m.group(1)) if m else None


# The four headings the author uses, with or without a trailing colon.
SECT = ["Early Game", "How to trade", "What to watch out for", "Tips"]


def split_sections(txt):
    """Split the writeup on the author's own headings. Text is never altered."""
    hits = []
    for s in SECT:
        m = re.search(r'(?:^|\n)[ \t]*' + re.escape(s) + r'[ \t]*:?[ \t]*(?=\n|$)', txt, re.I)
        if m:
            hits.append((m.start(), m.end(), s))
    if not hits:
        return [{"heading": None, "body": txt.strip()}]
    hits.sort()
    out = []
    pre = txt[:hits[0][0]].strip()
    if pre:
        out.append({"heading": None, "body": pre})
    for k, (a, b, s) in enumerate(hits):
        end = hits[k + 1][0] if k + 1 < len(hits) else len(txt)
        out.append({"heading": s, "body": txt[b:end].strip()})
    return out


# A line the author started as a list item: "- x", "1. x", "2) x".
LIST_START = re.compile(r'^(?:-\s|\d+[.)]\s)')


def bullets(body):
    """Group the body into paragraphs on the author's own blank lines and list
    markers. Soft line wrapping inside a cell is joined; nothing is split,
    merged across a blank line, or reworded.

    Numbered items matter: several writeups (Sylas, Vladimir) use "1. / 2. / 3."
    on adjacent lines with no blank between, and joining those collapses a real
    list into one long paragraph."""
    items, buf = [], []
    for ln in body.split("\n"):
        s = ln.strip()
        if LIST_START.match(s):
            if buf:
                items.append(" ".join(buf).strip())
            # Drop the author's "- " dash (the layout draws its own marker) but
            # keep "1." / "2)" — those numbers carry meaning in the prose.
            buf = [s[2:] if s.startswith("- ") else s]
        elif not s:
            if buf:
                items.append(" ".join(buf).strip())
            buf = []
        else:
            buf.append(s)
    if buf:
        items.append(" ".join(buf).strip())
    return [i for i in items if i]


def parse_runes(txt):
    groups = [p.strip() for p in re.split(r'\n\s*\n', txt) if p.strip()]

    def toks(p):
        return [t.strip(" ^") for t in re.split(r'-\s*>|\+', p.replace("\n", " ")) if t.strip(" ^")]
    return {"primary": toks(groups[0]) if len(groups) > 0 else [],
            "secondary": toks(groups[1]) if len(groups) > 1 else [],
            "shards": toks(groups[2]) if len(groups) > 2 else [],
            "raw": txt}


def parse_build_text(txt):
    """-> ordered [{item, note}] in the author's original sequence.

    A bracketed aside such as "[ Can Go Plated Steelcaps Rush ]" is kept where
    the author put it (flagged `aside`) rather than hoisted to the top."""
    out, cur = [], None
    for ln in txt.split("\n"):
        s = ln.strip()
        if not s or set(s) <= set("-–— "):
            continue
        if s.startswith("("):
            note = s.strip("()").strip()
            if cur:
                cur["note"] = (cur["note"] + " " + note).strip()
            else:
                out.append({"item": "", "note": note, "aside": True})
            continue
        if s.endswith(")") and "(" not in s:
            if cur:
                cur["note"] = (cur["note"] + " " + s.rstrip(")").strip()).strip()
            continue
        if s.startswith("[") or s.lower().startswith("if ") or s.endswith("then go:"):
            if cur:
                out.append(cur)
                cur = None
            out.append({"item": s.strip("[]").strip(), "note": "", "aside": True})
            continue
        if cur:
            out.append(cur)
        cur = {"item": s, "note": ""}
    if cur:
        out.append(cur)
    return out


def build_at(r0, off_lbl, off_icon, off_txt, champ):
    icons = [IMG.get((r0 + off_icon, c)) for c in range(8, 14)]
    txt = g(r0 - 1 + off_txt, 8)
    label = g(r0 - 1 + off_lbl, 8)
    if not any(icons) and not txt:
        return None
    used_media.update(m for m in icons if m)
    # "Items vs." and "Items vs. Darius" carry no information beyond the page
    # heading; anything else the author wrote is a real condition and is kept.
    generic = re.fullmatch(r'items\s+vs\.?\s*(?:' + re.escape(champ.lower()) + r')?\s*:?',
                           label.lower().strip())
    return {"condition": "" if generic else label,
            "icons": [{"file": m, "name": ITEM_NAMES.get(m), "placeholder": PLACEHOLDER.get(m)}
                      for m in icons if m],
            "steps": parse_build_text(txt), "raw": txt}


anchors = [i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("c:")]
matchups = []
for a in anchors:
    r0 = a + 1
    raw_name = g(a, 2)
    name = CANON.get(raw_name, raw_name.strip())

    # Most blocks put the writeup at offset 0. A few (Udyr) instead hold a short
    # variant label there — "AD Udyr" — and carry two complete sub-matchups, each
    # with its own ratings, writeup and build. Detect and keep both.
    head0, head1 = g(a, 18), g(a + 1, 18)
    variant_layout = len(head0) < 40 and len(head1) > 200

    def variant(off_r, off_txt, off_lbl, off_icon, off_bt):
        return {
            "label": g(a + off_r, 18) if variant_layout else "",
            "ratings": {"early": rating(g(a + off_r, 14)), "mid": rating(g(a + off_r, 15)),
                        "late": rating(g(a + off_r, 16)), "overall": rating(g(a + off_r, 17))},
            "gameplayRaw": g(a + off_txt, 18),
            "sections": [dict(s, items=bullets(s["body"]))
                         for s in split_sections(g(a + off_txt, 18))],
            "build": build_at(r0, off_lbl, off_icon, off_bt, name),
        }

    if variant_layout:
        variants = [variant(0, 1, 0, 1, 2), variant(3, 4, 3, 4, 5)]
    else:
        variants = [variant(0, 0, 0, 1, 2)]
        second = build_at(r0, 3, 4, 5, name)
        if second:
            variants[0]["extraBuilds"] = [second]

    builds = []
    for v in variants:
        if v["build"]:
            builds.append(v["build"])
        for b in v.get("extraBuilds", []) or []:
            builds.append(b)

    port, key = IMG.get((r0, 3)), IMG.get((r0, 5))
    s1, s2 = IMG.get((r0 + 1, 16)), IMG.get((r0 + 1, 14))
    used_media.update(x for x in (port, key, s1, s2) if x)

    video_txt = g(a, 25) or g(a, 26) or g(a, 27)
    video_url = link(MATCHUPS_SHEET, r0, 25) or link(MATCHUPS_SHEET, r0, 26) \
        or link(MATCHUPS_SHEET, r0, 27)

    base = variants[0]
    matchups.append({
        "name": name, "slug": slug(name), "sourceName": raw_name.strip(),
        "aliases": ALIAS.get(name, []),
        "portrait": port, "keystoneIcon": key,
        "summoners": ([{"file": s1, "name": SUMMONERS.get(s1)},
                       {"file": s2, "name": SUMMONERS.get(s2)}] if s1 else []),
        "ratings": base["ratings"],
        "keystone": (parse_runes(g(a + 4, 5))["primary"] or [None])[0],
        "runes": parse_runes(g(a + 4, 5)),
        "builds": builds,
        "sections": base["sections"], "gameplayRaw": base["gameplayRaw"],
        # Only populated when the block really holds two sub-matchups.
        "variants": ([{k: v[k] for k in ("label", "ratings", "sections", "gameplayRaw")}
                      for v in variants] if len(variants) > 1 else []),
        "tldr": bullets(g(a + 2, 14)) if g(a + 2, 14) else [],
        "video": video_txt, "videoUrl": video_url,
    })

# ---------------------------------------------------------------- introduction
S = "Introduction"
irows = grid(S, 25, 16)
intro = {"title": cell(irows, 1, 1),
         "subtitle": cell(irows, 2, 1), "subtitleUrl": link(S, 2, 1),
         "toc": [{"label": cell(irows, r, 3), "url": link(S, r, 3)}
                 for r in range(4, 10) if cell(irows, r, 3)],
         "betaNote": cell(irows, 3, 9), "todo": cell(irows, 4, 9),
         "wikiLabel": cell(irows, 10, 1), "wikiUrl": link(S, 10, 1),
         "buglist": cell(irows, 10, 9), "buglistUrl": link(S, 10, 9),
         "skinNote": cell(irows, 12, 9), "skinUrl": link(S, 12, 9)}
patches, cur = [], None
for r in range(12, 26):
    a, b = cell(irows, r, 1), cell(irows, r, 2)
    if re.match(r'^V\d+\.\d+$', a):
        if cur:
            patches.append(cur)
        cur = {"version": a, "changes": []}
    elif a.startswith("- ") and cur is not None:
        cur["changes"].append({"ability": a[2:], "text": ""})
    elif b and cur is not None and cur["changes"]:
        cur["changes"][-1]["text"] = b
if cur:
    patches.append(cur)
intro["patches"] = patches

# ---------------------------------------------------------------- item guide
S = "Itemization Guide"
grows = grid(S, 46, 8)
d3 = drawing_map(3)
SLOTC = {1: "First", 3: "Second", 5: "Third", 7: "Fourth+"}
itemguide, cursec = [], None
for r in range(4, 46):
    head = cell(grows, r, 1)
    if head and not d3.get((r, 1)) and head != "General\nItemization\nGuide:":
        if cursec:
            itemguide.append(cursec)
        cursec = {"title": head, "entries": []}
    slots = []
    for c, lab in SLOTC.items():
        med, txt = d3.get((r, c)), cell(grows, r, c + 1)
        if med or txt:
            # Each cell is its own item: row 7 holds Rylai's as a first item and
            # Cosmic Drive as a second. Collapsing to one row icon is wrong.
            slots.append({"slot": lab, "icon": med, "text": txt})
    if slots and cursec is not None:
        cursec["entries"].append({"row": r, "slots": slots})
        used_media.update(s["icon"] for s in slots if s["icon"])
if cursec:
    itemguide.append(cursec)

# ---------------------------------------------------------------- rune guide
S = "Rune Guide"
rrows = grid(S, 28, 8)
d4 = drawing_map(4)
runeguide, cur = [], None
for r in range(1, 29):
    if cell(rrows, r, 1) == "Rune":
        cur = {"row": r, "entries": []}
        runeguide.append(cur)
        continue
    if cur is None:
        continue
    for c in (1, 3, 5, 7):
        med, txt = d4.get((r, c)), cell(rrows, r, c + 1)
        # "Add mini runes here" is the author's own scaffolding, not content.
        if txt.lower().startswith("add mini runes"):
            txt = ""
        if med or txt:
            cur["entries"].append({"icon": med, "text": txt})
            if med:
                used_media.add(med)
runeguide = [g_ for g_ in runeguide if any(e["text"] or e["icon"] for e in g_["entries"])]

# ---------------------------------------------------------------- alt setups
S = "Alternative Mordekaiser Setups"
arows = grid(S, 17, 7)
alts, group = [], None
for r in range(4, 18):
    a = cell(arows, r, 1)
    if a.endswith("entries:"):
        group = a.replace(" entries:", "")
        continue
    if a and a != "Build Concept" and cell(arows, r, 3):
        alts.append({"group": group, "name": a, "runes": cell(arows, r, 3),
                     "items": cell(arows, r, 4), "example": cell(arows, r, 5),
                     "exampleUrl": link(S, r, 5), "comments": cell(arows, r, 7)})

# ---------------------------------------------------------------- references
S = "Mordekaiser Content  References"
crows = grid(S, 40, 11)


def person(r):
    return {"region": cell(crows, r, 1), "name": cell(crows, r, 2),
            "peak": cell(crows, r, 3),
            "twitter": cell(crows, r, 4), "twitterUrl": link(S, r, 4),
            "youtube": cell(crows, r, 5), "youtubeUrl": link(S, r, 5),
            "twitch": cell(crows, r, 6), "twitchUrl": link(S, r, 6),
            "alt": cell(crows, r, 7), "altUrl": link(S, r, 7),
            "opgg": cell(crows, r, 8), "opggUrl": link(S, r, 8)}


def resource(r):
    return {"region": cell(crows, r, 1), "name": cell(crows, r, 2),
            "peak": cell(crows, r, 3),
            "title": cell(crows, r, 4), "url": link(S, r, 4),
            "date": cell(crows, r, 7)}


# The sheet holds three tables with different shapes; find each by its own header.
sections, cur, kind = [], None, None
for r in range(1, 41):
    a = cell(crows, r, 1)
    if a and a not in ("Region",) and not cell(crows, r, 2) and len(a) > 12:
        if cur and cur["rows"]:
            sections.append(cur)
        cur = {"title": a, "rows": [], "kind": "people"}
        kind = None
        continue
    if a == "Region":
        kind = "resources" if cell(crows, r, 4).lower().startswith("title") else "people"
        if cur:
            cur["kind"] = kind
        continue
    if cur is not None and cell(crows, r, 2):
        cur["rows"].append(resource(r) if cur["kind"] == "resources" else person(r))
if cur and cur["rows"]:
    sections.append(cur)

notes = [{"text": cell(crows, r, 1), "url": link(S, r, 1)}
         for r in range(31, 41) if cell(crows, r, 1)]
credits = [{"name": cell(crows, r, 10), "role": cell(crows, r, 11)}
           for r in range(3, 40) if cell(crows, r, 10)]

# ---------------------------------------------------------------- emit
os.makedirs(f"{OUT}/assets/img", exist_ok=True)
for m in sorted(used_media):
    src = f"{UNZ}/xl/media/{m}"
    if os.path.exists(src):
        shutil.copy(src, f"{OUT}/assets/img/{m}")
os.makedirs(f"{OUT}/data", exist_ok=True)
json.dump(matchups, open(f"{OUT}/data/matchups.json", "w"), indent=0, ensure_ascii=False)
json.dump({"intro": intro, "itemguide": itemguide, "runeguide": runeguide, "alts": alts,
           "referenceSections": sections, "notes": notes, "credits": credits},
          open(f"{OUT}/data/guides.json", "w"), indent=0, ensure_ascii=False)

shutil.rmtree(UNZ, ignore_errors=True)
print(f"matchups: {len(matchups)} | media: {len(used_media)} | hyperlinks: {len(LINKS)}")
print("patches:", [p["version"] for p in patches])
print("reference sections:", [(s["title"][:34], s["kind"], len(s["rows"])) for s in sections])
print("variants:", [m["name"] for m in matchups if m["variants"]])
print("builds with a real condition label:",
      sum(1 for m in matchups for b in m["builds"] if b["condition"]))
